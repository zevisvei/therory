import csv
from pathlib import Path
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

csv_file_path = Path(__file__).parent / "theoryexamhe-data-yellow-05-12.csv"
xlsx_output_path = Path(__file__).parent / "theory.xlsx"
images_dir = Path(__file__).parent / "images"
images_dir.mkdir(exist_ok=True)

headers = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/58.0.3029.110 Safari/537.3"
}


def is_valid_url(url: str) -> bool:
    p = urlparse(url)
    return p.scheme in ("http", "https") and bool(p.netloc)


wb = Workbook()
ws = wb.active
ws.title = "theory"
# הגדרת כיוון מימין לשמאל
ws.sheet_view.rightToLeft = True

ws.append([
    "question_id",
    "question_text",
    "option 1",
    "option 2",
    "option 3",
    "option 4",
    "correct_answer",
    "question_type",
    "category",
    "image"
])

with csv_file_path.open("r", newline="", encoding="utf-8-sig") as f:
    reader = csv.reader(f)
    next(reader)

    row_idx = 2  # שורה ראשונה אחרי הכותרות

    for row in reader:
        question, answer, category = row
        question_id, question_text = question.split(".", 1)
        question_id_int = int(question_id.strip())
        question_text = question_text.strip()

        soup = BeautifulSoup(answer, "html.parser")

        options = soup.find("ul").find_all("li")
        extracted_options = []
        correct_answer = None

        for i, li in enumerate(options, start=1):
            span = li.find("span")
            if span.get("id") == f"correctAnswer{question_id}":
                correct_answer = i
            extracted_options.append(span.get_text(strip=True))

        type_span = soup.find("span", string=lambda s: s and "|" in s)
        question_type_list = [
            x.strip().replace("«", "").replace("»", "")
            for x in type_span.get_text().split("|")
            if x.strip()
        ]

        img_url = ""
        img_tag = soup.find("img")
        if img_tag:
            img_url = img_tag.get("src")

        ws.append([
            question_id_int,
            question_text,
            *extracted_options,
            correct_answer,
            ", ".join(question_type_list),
            category,
            ""  # עמודה לתמונה
        ])

        # הטמעת תמונה
        if img_url and is_valid_url(img_url):
            img_url = img_url.strip().strip('"').strip("'")
            img_path = images_dir / f"{question_id_int}.jpg"
            if not img_path.exists():
                print(img_url)
                r = requests.get(img_url, headers=headers, timeout=10, verify=False)
                content_type = r.headers.get("Content-Type", "")
                if not content_type.startswith("image/"):
                    print(f"⏭ לא תמונה ({content_type}): {img_url}")
                    continue
                img_path.write_bytes(r.content)

            img = XLImage(str(img_path))
            img.width = 200
            img.height = 150

            # התאמה של גובה שורה ורוחב עמודה
            col_letter = get_column_letter(10)  # J = 10
            ws.row_dimensions[row_idx].height = 115  # מתאים לגובה התמונה
            ws.column_dimensions[col_letter].width = 30  # מתאים לרוחב התמונה

            # הוספת תמונה
            ws.add_image(img, f"{col_letter}{row_idx}")

        row_idx += 1

wb.save(xlsx_output_path)
